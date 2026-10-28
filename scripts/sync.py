"""Sync the Atlas from Drive exports: Papers Master (.xlsx) -> data/papers.json, notes PDFs -> public/notes/<Key>.pdf.

The Google Sheet stays the single source of truth; this script only reads exports of it and never edits Drive.

usage:
  python scripts/sync.py --xlsx <Papers Master.xlsx> [--notes-dir <folder with downloaded notes PDFs>]
                         [--drive-listing <notes folder listing .json>]

What it does
  1. Reads the "Papers" tab (columns A-N; column O "Our notes" is private team scribble and is NOT published).
  2. Matches each notes PDF to a Key by first-author surname tokens + year (e.g. Allen_Zhu_2019_Study_Notes.pdf ->
     allenzhu2019convergence). Unmatched or ambiguous files are reported, never guessed.
  3. Copies matched notes to public/notes/<Key>.pdf and records them in data/notes_map.json.
  4. Writes data/papers.json, adding: short name, notes path, visual flag (papers/<Key>/visual.js exists),
     tags + result from papers/<Key>/meta.json, thumbnail path if public/thumbs/<Key>.webp exists.
"""
import argparse
import datetime as dt
import json
import os
import re
import shutil
import sys

import openpyxl

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
COLS = ["Key", "Title", "Authors", "Year", "Venue", "Phase", "Survey section", "Lead (auto)", "Priority", "Link", "PDF",
        "Summary (AI)", "Assumptions & limits (AI)", "Status"]


def first_surname(authors):
    a = re.sub(r"\bet al\.?", "", authors.split(",")[0]).strip()
    return a.split()[-1] if a else ""


def tokens(s):
    return re.findall(r"[a-z]+|\d+", s.lower())


def read_sheet(path):
    wb = openpyxl.load_workbook(path)
    ws = wb["Papers"]
    hdr = [c.value for c in ws[1]]
    missing = [c for c in COLS if c not in hdr]
    if missing:
        sys.exit(f"Sheet columns changed; missing {missing}. Update COLS in scripts/sync.py.")
    ix = {c: hdr.index(c) for c in COLS}
    rows = []
    for r in ws.iter_rows(min_row=2):
        key = r[ix["Key"]].value
        if not key:
            continue
        link_cell, pdf_cell = r[ix["Link"]], r[ix["PDF"]]
        rows.append({
            "key": str(key).strip(),
            "title": str(r[ix["Title"]].value or "").strip(),
            "authors": str(r[ix["Authors"]].value or "").strip(),
            "year": int(float(r[ix["Year"]].value)) if r[ix["Year"]].value else None,
            "venue": str(r[ix["Venue"]].value or "").strip(),
            "phase": str(r[ix["Phase"]].value or "").strip(),
            "section": str(r[ix["Survey section"]].value or "").strip(),
            "priority": str(r[ix["Priority"]].value or "").strip(),
            "link": link_cell.hyperlink.target if link_cell.hyperlink else None,
            "pdf": pdf_cell.hyperlink.target if pdf_cell.hyperlink else None,
            "summary": str(r[ix["Summary (AI)"]].value or "").replace("\r\n", "\n").strip(),
            "assumptions": str(r[ix["Assumptions & limits (AI)"]].value or "").replace("\r\n", "\n").strip(),
            "sheet_status": str(r[ix["Status"]].value or "").strip(),
        })
    return rows


def match_notes(files, rows):
    out, problems = {}, []
    for f in files:
        tk = set(tokens(os.path.splitext(os.path.basename(f))[0]))
        cands = [r for r in rows if str(r["year"]) in tk and all(p in tk for p in tokens(first_surname(r["authors"])))]
        if len(cands) == 1:
            out[f] = cands[0]["key"]
        else:
            problems.append((f, [c["key"] for c in cands]))
    return out, problems


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--notes-dir")
    ap.add_argument("--drive-listing", help="JSON list of {id,title} for the Drive notes folder (optional, recorded in notes_map)")
    a = ap.parse_args()
    rows = read_sheet(a.xlsx)
    keys = [r["key"] for r in rows]
    dup = {k for k in keys if keys.count(k) > 1}
    if dup:
        sys.exit(f"Duplicate Keys in the sheet: {dup}")
    map_path = os.path.join(ROOT, "data", "notes_map.json")
    nmap = json.load(open(map_path, encoding="utf-8")) if os.path.exists(map_path) else {}
    drive = {d["title"]: d["id"] for d in json.load(open(a.drive_listing, encoding="utf-8"))} if a.drive_listing else {}
    if a.notes_dir:
        files = [os.path.join(a.notes_dir, f) for f in os.listdir(a.notes_dir) if f.lower().endswith(".pdf")]
        matched, problems = match_notes(files, rows)
        os.makedirs(os.path.join(ROOT, "public", "notes"), exist_ok=True)
        for f, k in matched.items():
            shutil.copyfile(f, os.path.join(ROOT, "public", "notes", f"{k}.pdf"))
            nmap[k] = {"file": os.path.basename(f), "drive_id": drive.get(os.path.basename(f), nmap.get(k, {}).get("drive_id"))}
            print(f"notes  {os.path.basename(f):45s} -> {k}")
        for f, c in problems:
            print(f"UNMATCHED notes file {os.path.basename(f)} (candidates: {c or 'none'}) — rename it or add it to notes_map.json by hand")
        json.dump(nmap, open(map_path, "w", encoding="utf-8"), indent=1)
    papers = []
    for r in rows:
        k = r["key"]
        meta_p = os.path.join(ROOT, "papers", k, "meta.json")
        meta = json.load(open(meta_p, encoding="utf-8")) if os.path.exists(meta_p) else {}
        notes_file = os.path.join(ROOT, "public", "notes", f"{k}.pdf")
        sn = first_surname(r["authors"])
        papers.append({**r,
                       "short": f"{sn} {r['year']}",
                       "notes": f"public/notes/{k}.pdf" if os.path.exists(notes_file) else None,
                       "notes_title": nmap.get(k, {}).get("file"),
                       "visual": os.path.exists(os.path.join(ROOT, "papers", k, "visual.js")),
                       "tags": meta.get("tags", []),
                       "result": meta.get("result_plain") or "",
                       "tabs": meta.get("tabs", []),
                       "tag_tabs": meta.get("tag_tabs", {}),
                       "thumb": next((f"public/thumbs/{k}.{e}" for e in ("webp", "png") if os.path.exists(os.path.join(ROOT, "public", "thumbs", f"{k}.{e}"))), None)})
    out = {"synced": dt.datetime.now().strftime("%Y-%m-%d %H:%M"), "source": "Google Sheet 'Papers Master' (Drive: MLR570 Project / NonConvex Survey)", "papers": papers}
    json.dump(out, open(os.path.join(ROOT, "data", "papers.json"), "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    n_notes = sum(1 for p in papers if p["notes"])
    n_vis = sum(1 for p in papers if p["visual"])
    print(f"papers.json: {len(papers)} papers, {n_notes} with notes, {n_vis} with visuals")
    no_notes_but_read = [p["key"] for p in papers if p["sheet_status"] == "Read" and not p["notes"]]
    if no_notes_but_read:
        print("marked Read in the sheet but no notes PDF:", no_notes_but_read)
    notes_no_visual = [p["key"] for p in papers if p["notes"] and not p["visual"]]
    if notes_no_visual:
        print("NEW WORK — notes but no visual yet:", notes_no_visual)


if __name__ == "__main__":
    main()
